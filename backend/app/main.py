import hashlib
import logging
import os
import sys
# 设置环境变量确保UTF-8编码
os.environ['PYTHONIOENCODING'] = 'utf-8'

# 确保标准输出流使用 UTF-8 编码
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

# 设置默认编码为UTF-8
if sys.version_info[0] == 3 and sys.version_info[1] >= 7:
    # 对于Python 3.7及以上版本
    import io
    sys.stdout = io.TextIOWrapper(
        sys.stdout.buffer,
        encoding='utf-8',
        errors='replace',
        line_buffering=True
    )
    sys.stderr = io.TextIOWrapper(
        sys.stderr.buffer,
        encoding='utf-8',
        errors='replace',
        line_buffering=True
    )

import base64
import json
import math
import shutil
import io
import mimetypes
import tempfile
import threading
import time
import zipfile
from collections import OrderedDict
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from typing import Iterable, List, Mapping

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches, Cm
import cv2
from fastapi import FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse

try:
    from .ctc import CTCAnalyzer
except ImportError:
    # 当直接运行脚本时使用绝对导入
    from ctc import CTCAnalyzer

for _stream_name in ("stdout", "stderr"):
    _stream = getattr(sys, _stream_name, None)
    if _stream is None:
        continue
    encoding = getattr(_stream, "encoding", None)
    if encoding and encoding.lower() == "utf-8":
        continue
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except AttributeError:
        buffer = getattr(_stream, "buffer", None)
        if buffer is not None:
            setattr(sys, _stream_name, io.TextIOWrapper(buffer, encoding="utf-8", errors="replace"))

LOG_DIR = Path(__file__).resolve().parent.parent / "log"
LOG_DIR.mkdir(parents=True, exist_ok=True)
LOG_FILE = LOG_DIR / "app.log"

_file_handler = logging.FileHandler(LOG_FILE, encoding="utf-8")
_stream_handler = logging.StreamHandler(sys.stdout)
logging.basicConfig(
    level=logging.INFO,
    format="[%(asctime)s] %(levelname)s %(name)s - %(message)s",
    handlers=[_file_handler, _stream_handler],
)

logger = logging.getLogger("ctc_app")

INLINE_SUPPORTED_MIME_TYPES = {"image/png", "image/jpeg", "image/gif"}
MASK_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}

SONGTI_FONT_NAME = "SimSun"
TITLE_FONT_SIZE_PT = 18
BODY_FONT_SIZE_PT = 11

CHANNEL_SIGNAL_SUFFIXES = ("绿色信号", "红色信号", "蓝色信号")
CHANNEL_SUFFIX_LABEL_MAP = {
    "g": "绿色信号",
    "r": "红色信号",
    "b": "蓝色信号",
}

SAMPLE_NUMBER_STORAGE_DIR = Path(__file__).resolve().parent.parent / "var"
SAMPLE_NUMBER_STORAGE_DIR.mkdir(parents=True, exist_ok=True)
SAMPLE_NUMBER_STORAGE_PATH = SAMPLE_NUMBER_STORAGE_DIR / "sample_number_sequence.json"
REPORT_NUMBER_STORAGE_PATH = SAMPLE_NUMBER_STORAGE_DIR / "report_number_sequence.json"

SAMPLE_NUMBER_LOCK = threading.Lock()
REPORT_NUMBER_LOCK = threading.Lock()

SAMPLE_NUMBER_DIGITS = 7
REPORT_NUMBER_PREFIX = "PE"

CAT_ALIASES = {"猫", "貓", "cat"}
DOG_ALIASES = {"狗", "犬", "dog"}
OTHER_ALIASES = {"other", "others", "其他", "其它"}
CAT_ALIAS_KEYS = {alias.casefold() for alias in CAT_ALIASES}
DOG_ALIAS_KEYS = {alias.casefold() for alias in DOG_ALIASES}
OTHER_ALIAS_KEYS = {alias.casefold() for alias in OTHER_ALIASES}
DEFAULT_SAMPLE_NUMBER_PREFIX = "OTH"

DISCLAIMER_LINES = [
    "本检测结果仅供科研及临床辅助参考，不能作为唯一诊断依据。",
    "建议结合兽医临床表现、影像学及其他实验室检查综合判断。",
    "检测人____________                审核人____________                签字/盖章____________",
]


def _normalize_mask_source_name(value: str | None) -> str:
    if not value:
        return ""
    normalized = str(value).strip()
    if not normalized:
        return ""
    normalized = normalized.replace("\\", "/")
    last_segment = normalized.split("/")[-1]
    label_split = last_segment.split("：")[-1]
    colon_split = label_split.split(":")[-1]
    return colon_split.strip().casefold()


def _detect_channel_suffix(value: str | None) -> str | None:
    normalized = _normalize_mask_source_name(value)
    if not normalized:
        return None
    normalized = normalized.replace("_overlay", "")
    if "." in normalized:
        normalized = normalized.rsplit(".", 1)[0]
    normalized = normalized.strip()
    if not normalized:
        return None
    suffix = normalized[-1]
    if suffix in CHANNEL_SUFFIX_LABEL_MAP:
        return suffix
    return None


def _resolve_channel_label_from_suffix(
    prefix: str, suffix: str | None, fallback_label: str
) -> str:
    if suffix and suffix in CHANNEL_SUFFIX_LABEL_MAP:
        return f"{prefix}{CHANNEL_SUFFIX_LABEL_MAP[suffix]}"
    return fallback_label


def _normalize_pet_type(value: str | None) -> str:
    if not value:
        return ""
    return value.strip()


def _resolve_sample_number_prefix(pet_type: str | None) -> str:
    if not pet_type:
        return DEFAULT_SAMPLE_NUMBER_PREFIX

    normalized = pet_type.strip().casefold()
    if not normalized:
        return DEFAULT_SAMPLE_NUMBER_PREFIX

    if normalized in CAT_ALIAS_KEYS:
        return "CAT"
    if normalized in DOG_ALIAS_KEYS:
        return "DOG"
    if normalized in OTHER_ALIAS_KEYS:
        return DEFAULT_SAMPLE_NUMBER_PREFIX
    return DEFAULT_SAMPLE_NUMBER_PREFIX


def _load_number_counters(storage_path: Path, description: str) -> dict[str, int]:
    if not storage_path.exists():
        return {}
    try:
        with storage_path.open("r", encoding="utf-8") as file_obj:
            data = json.load(file_obj)
    except (OSError, json.JSONDecodeError):
        logger.warning("无法读取%s序列文件，将重新初始化：%s", description, storage_path)
        return {}

    counters: dict[str, int] = {}
    if isinstance(data, dict):
        for prefix, value in data.items():
            try:
                counters[str(prefix)] = int(value)
            except (TypeError, ValueError):
                logger.debug("忽略无效的%s序列值：prefix=%s value=%s", description, prefix, value)
                continue
    return counters


def _save_number_counters(
    storage_path: Path, counters: Mapping[str, int], description: str
) -> None:
    try:
        with storage_path.open("w", encoding="utf-8") as file_obj:
            json.dump(counters, file_obj, ensure_ascii=False, indent=2)
    except OSError:
        logger.exception("保存%s序列文件失败：%s", description, storage_path)


def _generate_next_counter_value(
    prefix: str,
    *,
    storage_path: Path,
    lock: threading.Lock,
    description: str,
) -> int:
    with lock:
        counters = _load_number_counters(storage_path, description)
        next_value = counters.get(prefix, -1) + 1
        counters[prefix] = next_value
        _save_number_counters(storage_path, counters, description)
    return next_value


def _sync_counter_value(
    prefix: str,
    numeric_value: int,
    *,
    storage_path: Path,
    lock: threading.Lock,
    description: str,
) -> None:
    with lock:
        counters = _load_number_counters(storage_path, description)
        current_max = counters.get(prefix, -1)
        if numeric_value > current_max:
            counters[prefix] = numeric_value
            _save_number_counters(storage_path, counters, description)


def _generate_next_sample_number(pet_type: str | None) -> str:
    prefix = _resolve_sample_number_prefix(pet_type)
    next_value = _generate_next_counter_value(
        prefix,
        storage_path=SAMPLE_NUMBER_STORAGE_PATH,
        lock=SAMPLE_NUMBER_LOCK,
        description="样本编号",
    )
    sample_number = f"{prefix}{next_value:0{SAMPLE_NUMBER_DIGITS}d}"
    logger.info("生成样本编号：%s（宠物类型=%s）", sample_number, pet_type or "未填写")
    return sample_number


def _sync_sample_number_counter(prefix: str, numeric_value: int) -> None:
    _sync_counter_value(
        prefix,
        numeric_value,
        storage_path=SAMPLE_NUMBER_STORAGE_PATH,
        lock=SAMPLE_NUMBER_LOCK,
        description="样本编号",
    )


def _generate_next_report_number() -> str:
    next_value = _generate_next_counter_value(
        REPORT_NUMBER_PREFIX,
        storage_path=REPORT_NUMBER_STORAGE_PATH,
        lock=REPORT_NUMBER_LOCK,
        description="报告编号",
    )
    report_number = f"{REPORT_NUMBER_PREFIX}{next_value:0{SAMPLE_NUMBER_DIGITS}d}"
    logger.info("生成报告编号：%s", report_number)
    return report_number


def _sync_report_number_counter(numeric_value: int) -> None:
    _sync_counter_value(
        REPORT_NUMBER_PREFIX,
        numeric_value,
        storage_path=REPORT_NUMBER_STORAGE_PATH,
        lock=REPORT_NUMBER_LOCK,
        description="报告编号",
    )


def verify_startup_token() -> None:
    """验证启动token，确保只能通过前端启动"""
    token = os.environ.get("FASTAPI_STARTUP_TOKEN")
    if not token:
        logger.error("错误：缺少启动token，此应用只能通过前端启动")
        sys.exit(1)

    try:
        timestamp_str, hash_part = token.split("_", 1)
        timestamp = int(timestamp_str)
        current_time = int(time.time())

        if current_time - timestamp > 60:
            logger.error("错误：启动token已过期")
            sys.exit(1)

        expected_hash = hashlib.sha256(f"fastapi_startup_{timestamp}".encode()).hexdigest()[:16]
        if hash_part != expected_hash:
            logger.error("错误：无效的启动token")
            sys.exit(1)

        logger.info("启动token验证成功")
    except (ValueError, IndexError):
        logger.error("错误：启动token格式无效")
        sys.exit(1)


def _safe_extract(zip_file: zipfile.ZipFile, target_dir: Path) -> None:
    target_dir.mkdir(parents=True, exist_ok=True)
    target_root = target_dir.resolve()
    for member in zip_file.infolist():
        member_path = target_root / member.filename
        try:
            member_path.resolve().relative_to(target_root)
        except ValueError as exc:  # pragma: no cover - 安全检查
            raise HTTPException(status_code=400, detail="压缩包包含非法路径，无法解压") from exc
    zip_file.extractall(target_dir)


def _find_dataset_root(extracted_dir: Path) -> Path:
    expected_dirs = {str(i) for i in range(1, 6)}
    extracted_dir.mkdir(parents=True, exist_ok=True)

    def has_expected_children(path: Path) -> bool:
        return any((path / child).exists() for child in expected_dirs)

    if has_expected_children(extracted_dir):
        return extracted_dir

    candidates = [p for p in extracted_dir.iterdir() if p.is_dir()]
    for candidate in candidates:
        if has_expected_children(candidate):
            return candidate

    if len(candidates) == 1:
        return candidates[0]

    raise HTTPException(status_code=400, detail="未找到符合CTC分析目录结构的文件夹")


def _normalize_field(value: str) -> str:
    return value.strip()


def _create_metadata_entries(
    institution_name: str,
    report_number: str,
    detection_date: str,
    sample_number: str,
    sample_volume: str,
    sample_status: str,
    pet_type: str,
    cancer_biomarker: str,
    department: str,
    pet_name: str,
    owner_name: str,
    age: str,
    gender: str,
    sample_type: str,
    medication_intake: str,
    notes: str,
) -> OrderedDict:
    entries: OrderedDict[str, str] = OrderedDict()
    # 按照新顺序调整字段
    entries["宠主姓名"] = _normalize_field(owner_name)
    entries["宠物姓名"] = _normalize_field(pet_name)
    entries["性别"] = _normalize_field(gender)
    entries["宠物类型"] = _normalize_field(pet_type)
    entries["年龄"] = _normalize_field(age)
    entries["机构名称"] = _normalize_field(institution_name)
    entries["检测日期"] = _normalize_field(detection_date)
    entries["报告编号"] = _normalize_field(report_number).upper()
    entries["科别"] = _normalize_field(department)
    entries["癌症标志物"] = _normalize_field(cancer_biomarker)
    entries["标本类型"] = _normalize_field(sample_type)
    entries["样品量（单位ml）"] = _normalize_field(sample_volume)
    entries["样本编号"] = _normalize_field(sample_number).upper()
    entries["样本状态"] = _normalize_field(sample_status)
    entries["一周内是否有药物摄入"] = _normalize_field(medication_intake)
    entries["备注"] = _normalize_field(notes)
    return entries


def _encode_preview_image(image_path: Path) -> tuple[str, str] | None:
    """读取图像并转换为可预览的Base64格式"""
    try:
        mime_type, _ = mimetypes.guess_type(str(image_path))
        if mime_type in INLINE_SUPPORTED_MIME_TYPES:
            image_bytes = image_path.read_bytes()
            return mime_type, base64.b64encode(image_bytes).decode("ascii")

        image = cv2.imread(str(image_path), cv2.IMREAD_UNCHANGED)
        if image is None:
            image_bytes = image_path.read_bytes()
            return (mime_type or "image/png", base64.b64encode(image_bytes).decode("ascii"))

        success, buffer = cv2.imencode(".png", image)
        if not success:
            logger.warning("图像编码为PNG失败：%s", image_path)
            image_bytes = image_path.read_bytes()
            return (mime_type or "image/png", base64.b64encode(image_bytes).decode("ascii"))

        return "image/png", base64.b64encode(buffer.tobytes()).decode("ascii")
    except OSError:
        logger.warning("无法读取预览图像：%s", image_path)
        return None
    except Exception:
        logger.exception("处理预览图像时出现错误：%s", image_path)
        return None


def _apply_run_style(run, size: int, bold: bool = False, color: RGBColor | None = None) -> None:
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.name = SONGTI_FONT_NAME
    if color is not None:
        font.color.rgb = color

    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), SONGTI_FONT_NAME)
    r_fonts.set(qn("w:hAnsi"), SONGTI_FONT_NAME)
    r_fonts.set(qn("w:eastAsia"), SONGTI_FONT_NAME)
    r_fonts.set(qn("w:cs"), SONGTI_FONT_NAME)
    r_fonts.set(qn("w:hint"), "eastAsia")

    latin_font = r_pr.find(qn("w:latin"))
    if latin_font is None:
        latin_font = OxmlElement("w:latin")
        r_pr.append(latin_font)
    latin_font.set(qn("w:typeface"), SONGTI_FONT_NAME)


def _apply_font_size(paragraphs: Iterable, *, size: int = BODY_FONT_SIZE_PT) -> None:
    for paragraph in paragraphs:
        for run in paragraph.runs:
            _apply_run_style(run, size)


def _set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: RGBColor | None = None,
    fill: str | None = None,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    lines = text.splitlines() or [""]
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        _apply_run_style(run, BODY_FONT_SIZE_PT, bold=bold, color=color)
        if index < len(lines) - 1:
            run.add_break()
    paragraph.paragraph_format.space_after = Pt(0)

    if fill is not None:
        tc_pr = cell._tc.get_or_add_tcPr()
        for child in list(tc_pr):
            if child.tag == qn("w:shd"):
                tc_pr.remove(child)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), fill)
        tc_pr.append(shading)


def _set_label_value_cell(
    cell,
    label: str,
    value: str,
    *,
    label_color: RGBColor,
    value_color: RGBColor,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    label_run = paragraph.add_run(f"{label}：")
    _apply_run_style(label_run, BODY_FONT_SIZE_PT, bold=True, color=label_color)
    value_run = paragraph.add_run(value)
    _apply_run_style(value_run, BODY_FONT_SIZE_PT, color=value_color)
    paragraph.paragraph_format.space_after = Pt(0)


def _add_section_heading(document: Document, text: str, *, color: RGBColor) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    _apply_run_style(run, 16, bold=True, color=color)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)


def _populate_table(table, rows: List[List[str]]) -> None:
    for row_idx, row_values in enumerate(rows):
        row = table.rows[row_idx]
        for cell, value in zip(row.cells, row_values):
            cell.text = value
            _apply_font_size(cell.paragraphs)


def _set_table_transparent(table) -> None:
    tbl = table._tbl
    # 使用xpath获取tblPr，如果不存在则创建
    tblPr_elements = tbl.xpath('w:tblPr')
    if tblPr_elements:
        tblPr = tblPr_elements[0]
    else:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)

    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{border_name}"))
        if border is None:
            border = OxmlElement(f"w:{border_name}")
            borders.append(border)
        border.set(qn("w:val"), "nil")

    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()

            shading = tc_pr.find(qn("w:shd"))
            if shading is not None:
                tc_pr.remove(shading)

            tc_borders = tc_pr.find(qn("w:tcBorders"))
            if tc_borders is None:
                tc_borders = OxmlElement("w:tcBorders")
                tc_pr.append(tc_borders)

            for border_name in ("top", "left", "bottom", "right"):
                border = tc_borders.find(qn(f"w:{border_name}"))
                if border is None:
                    border = OxmlElement(f"w:{border_name}")
                    tc_borders.append(border)
                border.set(qn("w:val"), "nil")


def _format_channel_summary_texts(analyzer: CTCAnalyzer) -> list[str]:
    """Return a simplified summary for channel data."""

    # 需求更新：报告中不再展示原始通道统计数组
    return []


def _generate_result_description(
    metadata: Mapping[str, str] | None,
    ctc_counts: Iterable[int] | None,
    wbc_counts: Iterable[int] | None,
) -> tuple[str, list[str]]:
    ctc_list = list(ctc_counts or [])
    wbc_list = list(wbc_counts or [])

    if not ctc_list and not wbc_list:
        return "结果说明：", ["未识别出有效的检测结果，请检查上传的影像资料。"]

    normalized_metadata = metadata or {}
    biomarker_value = (normalized_metadata.get("癌症标志物") or "").strip()
    biomarker_label = biomarker_value if biomarker_value else "癌症标志物"

    ctc_first_two = sum(ctc_list[:2])
    wbc_first_two = sum(wbc_list[:2])

    sample_volume_raw = (normalized_metadata.get("样品量（单位ml）") or "").strip()
    sample_volume_ml = _parse_sample_volume_to_ml(sample_volume_raw)
    ctc_total = sum(ctc_list)
    wbc_total = sum(wbc_list)

    if sample_volume_ml is not None:
        ctc_cells_per_ml = _format_cells_per_ml(ctc_total, sample_volume_ml)
        wbc_cells_per_ml = _format_cells_per_ml(wbc_total, sample_volume_ml)
        ratio_line = (
            f"检测CTC数量：{ctc_cells_per_ml} cells/ml；"
            f"背景白细胞：{wbc_cells_per_ml} cells/ml"
        )
    else:
        sample_volume_text = sample_volume_raw if sample_volume_raw else "未填写"
        ratio_line = (
            f"检测CTC数量：{ctc_total}/{sample_volume_text} cells/ml；"
            f"背景白细胞：{wbc_total}/{sample_volume_text} cells/ml"
        )

    first_line = (
        f"经检测结果判定，在一二通道中找到CD45+ {wbc_first_two}个，"
        f"{biomarker_label}：{ctc_first_two}个。"
    )

    return "结果说明：", [first_line, ratio_line]


def _format_detection_result_rows(
    doc_names: Iterable[str],
    ctc_counts: Iterable[int],
    wbc_counts: Iterable[int],
) -> list[list[str]]:
    triplets = list(zip(doc_names, ctc_counts, wbc_counts))[:3]
    if not triplets:
        return []

    display_triplets: list[tuple[str, int, int]] = []
    entries: list[str] = []
    for index, (name, ctc, wbc) in enumerate(triplets):
        normalized_name = (name or "").strip()
        label = normalized_name if normalized_name else f"检测结果{index + 1}"
        if wbc == 0:
            continue

        display_triplets.append((label, ctc, wbc))
        entries.append(f"{label}: CD45 {wbc}个，CK {ctc}个")

    if not display_triplets:
        return []

    combined_wbc = sum(wbc for _, _, wbc in display_triplets)
    combined_ctc = sum(ctc for _, ctc, _ in display_triplets)
    combined_entry = f"合计: CD45 {combined_wbc}个，CK {combined_ctc}个"

    rows: list[list[str]] = []
    first_row_items = entries[:2]
    if first_row_items:
        rows.append(first_row_items)

    remaining_items = entries[2:]
    second_row_items = remaining_items + [combined_entry]
    if second_row_items:
        rows.append(second_row_items)

    return rows


def _format_channel_label_prefix(sub_folder: str | None) -> str:
    """Format the label prefix for channel previews using folder information."""

    normalized = (sub_folder or "").strip()
    if not normalized:
        return "X通道"
    if normalized.endswith("通道"):
        return normalized
    return f"{normalized}通道"


def _wrap_parentheses(value: str | None) -> str:
    if not value:
        return "（未填写）"
    normalized = value.strip()
    if not normalized:
        return "（未填写）"
    if normalized.startswith("（") and normalized.endswith("）"):
        return normalized
    return f"（{normalized}）"


def _wrap_ascii_parentheses(value: str | None, *, fallback: str = "未填写") -> str:
    return _ensure_value(value, fallback)


def _format_checkbox_line(label: str, value: str | None, options: list[str]) -> str:
    normalized_value = _ensure_value(value, "").strip()
    formatted_options: list[str] = []
    for option in options:
        mark = "√" if normalized_value == option else "¨"
        formatted_options.append(f"{option}{mark}")
    options_text = " ".join(formatted_options) if formatted_options else ""
    return f"{label}:  {options_text}".rstrip()


def _ensure_value(value: str | None, fallback: str = "未填写") -> str:
    if not value:
        return fallback
    normalized = value.strip()
    return normalized if normalized else fallback


def _parse_sample_volume_to_ml(value: str | None) -> float | None:
    if not value:
        return None
    normalized = value.strip().lower()
    if not normalized:
        return None
    compact = normalized.replace("毫升", "").replace("ml", "").replace(" ", "")
    compact = compact.replace(",", "")
    if not compact:
        return None
    try:
        volume = float(compact)
    except ValueError:
        logger.debug("无法解析样品量为浮点数：%s", value)
        return None
    if volume <= 0:
        logger.debug("样品量非正数，忽略：%s", value)
        return None
    return volume


def _format_cells_per_ml(total_cells: int, sample_volume_ml: float | None) -> str:
    if sample_volume_ml:
        value = total_cells / sample_volume_ml
        formatted = f"{value:.2f}".rstrip("0").rstrip(".")
        return formatted
    return str(total_cells)


def _format_sample_volume(value: str | None) -> str:
    normalized = _ensure_value(value)
    if normalized == "未填写":
        return normalized
    compact = normalized.replace(" ", "")
    if compact.lower().endswith("ml"):
        return normalized
    return f"{normalized} ml"


def _format_report_date(value: str | None) -> str:
    if not value:
        return ""
    normalized = value.strip()
    if not normalized:
        return ""
    parts = normalized.split("-")
    if len(parts) != 3 or not all(part.isdigit() for part in parts):
        return normalized
    year, month, day = parts
    return f"{year}年{int(month)}月{int(day)}日"


def _build_report_document(
    analyzer: CTCAnalyzer,
    metadata: OrderedDict[str, str],
    output_path: Path,
    channel_summary_texts: list[str] | None = None,
    mask_option: dict | None = None,
    *,
    mask_options: list[dict] | None = None,
    detection_result_rows: list[list[str]] | None = None,
    result_description: tuple[str, list[str]] | None = None,
) -> None:
    document = Document()

    for section in document.sections:
        section.top_margin = Cm(1.27)
        section.bottom_margin = Cm(1.27)
        section.left_margin = Cm(1.27)
        section.right_margin = Cm(1.27)
        section.gutter = Cm(0)
    # 修改logo路径查找逻辑，使其在打包后也能正确找到logo文件
    logo_path = Path(__file__).resolve().parent.parent / "HBI.png"
    
    # 如果在默认位置找不到logo，则尝试在exe文件同目录下查找
    if not logo_path.exists():
        if getattr(sys, 'frozen', False):
            # 如果是打包后的exe文件，尝试在exe同目录下查找
            exe_dir = Path(sys.executable).parent
            logo_path = exe_dir / "HBI.png"
    
    if logo_path.exists():
        logo_paragraph = document.add_paragraph()
        logo_run = logo_paragraph.add_run()
        logo_run.add_picture(str(logo_path), width=Inches(1.6))
        logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        logo_paragraph.paragraph_format.space_after = Pt(6)


    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand_run = title_paragraph.add_run("FlowCanis ")
    _apply_run_style(brand_run, TITLE_FONT_SIZE_PT, bold=True, color=RGBColor(0, 0, 0))
    main_title_run = title_paragraph.add_run("微流控循环肿瘤细胞分选")
    _apply_run_style(main_title_run, TITLE_FONT_SIZE_PT, bold=True, color=RGBColor(31, 41, 55))

    subtitle_paragraph = document.add_paragraph()
    subtitle_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_paragraph.add_run("与免疫荧光识别检测报告单")
    _apply_run_style(subtitle_run, TITLE_FONT_SIZE_PT, bold=True, color=RGBColor(31, 41, 55))

    report_number_paragraph = document.add_paragraph()
    report_number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    report_number_text = _ensure_value(metadata.get("报告编号") if metadata else None, "未填写")
    report_number_run = report_number_paragraph.add_run(f"编号：[{report_number_text}]")
    _apply_run_style(report_number_run, BODY_FONT_SIZE_PT, bold=True, color=RGBColor(0, 0, 0))  # 改为黑色
    report_number_paragraph.paragraph_format.space_after = Pt(12)

    if metadata:
        detection_date = _format_report_date(metadata.get("检测日期"))
        first_line_parts = [
            f"宠主姓名: {_wrap_ascii_parentheses(metadata.get('宠主姓名'))}",
            f"宠物姓名: {_wrap_ascii_parentheses(metadata.get('宠物姓名'))}",
            f"性别: {_wrap_ascii_parentheses(metadata.get('性别'))}",
            f"宠物类型: {_wrap_ascii_parentheses(metadata.get('宠物类型'))}",
            f"年龄: {_wrap_ascii_parentheses(metadata.get('年龄'))}",
        ]
        first_line = "   ".join(first_line_parts)

        second_line_parts = [
            f"送检单位: {_wrap_ascii_parentheses(metadata.get('机构名称'))}",
            f"送检时间: {_wrap_ascii_parentheses(detection_date or None)}",
            f"科别: {_wrap_ascii_parentheses(metadata.get('科别'))}",
            f"癌症标志物: {_wrap_ascii_parentheses(metadata.get('癌症标志物'))}",
        ]
        second_line = "   ".join(second_line_parts)

        formatted_sample_volume = _format_sample_volume(metadata.get("样品量（单位ml）"))
        if formatted_sample_volume != "未填写":
            formatted_sample_volume = formatted_sample_volume.replace(" ", "").upper()

        third_line = "   ".join(
            [
                f"样本类型: {_wrap_ascii_parentheses(metadata.get('标本类型'))}",
                f"样品量: {_wrap_ascii_parentheses(formatted_sample_volume)}",
                f"样本编号: {_wrap_ascii_parentheses(metadata.get('样本编号'))}",
                f"样本状态: {_wrap_ascii_parentheses(metadata.get('样本状态'))}",
            ]
        )

        medication_value = _wrap_ascii_parentheses(
            metadata.get("一周内是否有药物摄入")
        )
        notes_value = _wrap_ascii_parentheses(metadata.get("备注"), fallback="无")
        fourth_line = "   ".join(
            [
                f"一周内是否有药物摄入: {medication_value}",
                f"如有，请说明: {notes_value}",
            ]
        )

        divider_text = "─────────────────────────────────────────────────"
        info_lines = [first_line, second_line, third_line, fourth_line]

        info_paragraphs = []
        divider_paragraph = document.add_paragraph()
        divider_run = divider_paragraph.add_run(divider_text)
        _apply_run_style(divider_run, BODY_FONT_SIZE_PT, color=RGBColor(0, 0, 0))
        divider_paragraph.paragraph_format.space_after = Pt(0)
        for line in info_lines:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            _apply_run_style(run, BODY_FONT_SIZE_PT, color=RGBColor(31, 41, 55))
            paragraph.paragraph_format.space_after = Pt(0)
            info_paragraphs.append(paragraph)

        if info_paragraphs:
            info_paragraphs[-1].paragraph_format.space_after = Pt(0)

        bottom_divider = document.add_paragraph()
        bottom_run = bottom_divider.add_run(divider_text)
        _apply_run_style(bottom_run, BODY_FONT_SIZE_PT, color=RGBColor(0, 0, 0))
        bottom_divider.paragraph_format.space_after = Pt(6)


    result_heading = document.add_paragraph()
    result_heading_run = result_heading.add_run("检测结果")
    _apply_run_style(result_heading_run, 16, bold=True, color=RGBColor(37, 99, 235))
    result_heading.paragraph_format.space_before = Pt(12)
    result_heading.paragraph_format.space_after = Pt(6)

    doc_names = analyzer.results.get("doc_names", [])
    ctc_counts = analyzer.results.get("green_single_channel", [])
    wbc_counts = analyzer.results.get("white_single_channel", [])

    total_ctc = sum(ctc_counts)
    total_wbc = sum(wbc_counts)

    if detection_result_rows is None:
        detection_result_rows = _format_detection_result_rows(
            doc_names,
            ctc_counts,
            wbc_counts,
        )

    result_detail_paragraphs: list = []
    for row_entries in detection_result_rows:
        if not row_entries:
            continue
        detail_paragraph = document.add_paragraph()
        detail_run = detail_paragraph.add_run("   ".join(row_entries))
        _apply_run_style(detail_run, BODY_FONT_SIZE_PT, color=RGBColor(55, 65, 81))
        detail_paragraph.paragraph_format.space_after = Pt(0)
        result_detail_paragraphs.append(detail_paragraph)

    if result_detail_paragraphs:
        result_detail_paragraphs[-1].paragraph_format.space_after = Pt(4)

    ctc_image_sets = analyzer.results.get("ctc_image_sets", [])
    channel_label_prefix = _format_channel_label_prefix(None)
    if ctc_image_sets:
        image_set = ctc_image_sets[0]
        channel_label_prefix = _format_channel_label_prefix(
            image_set.get("sub_folder")
        )
        channel_signal_labels = [
            f"{channel_label_prefix}{suffix}"
            for suffix in CHANNEL_SIGNAL_SUFFIXES
        ]

        for summary_text in channel_summary_texts or []:
            if not summary_text:
                continue
            summary_paragraph = document.add_paragraph()
            summary_paragraph.paragraph_format.space_after = Pt(2)
            summary_run = summary_paragraph.add_run(summary_text)
            _apply_run_style(summary_run, BODY_FONT_SIZE_PT, color=RGBColor(55, 65, 81))

        def _resolve_preview_path(primary_key: str, *fallback_keys: str) -> str | None:
            keys = (primary_key,) + fallback_keys
            for key in keys:
                candidate = image_set.get(key)
                if candidate and os.path.exists(candidate):
                    return candidate
            return None

        selected_labels_and_paths: list[tuple[str, str, str | None]] = []

        if mask_options:
            for index, option in enumerate(mask_options):
                candidate_path = option.get("path")
                if not candidate_path or not os.path.exists(candidate_path):
                    logger.warning(
                        "所选掩码图像不存在或不可访问：%s", candidate_path
                    )
                    continue

                detail_label = option.get("label") or option.get("channel") or "掩码图像"
                normalized_detail = str(detail_label).strip()
                base_label = f"图像{index + 1}"
                if normalized_detail:
                    label = f"{base_label}：{normalized_detail}"
                else:
                    label = base_label
                color_suffix = (
                    _detect_channel_suffix(option.get("label"))
                    or _detect_channel_suffix(option.get("path"))
                    or _detect_channel_suffix(option.get("relativePath"))
                )
                selected_labels_and_paths.append((label, candidate_path, color_suffix))

        mask_override_path: str | None = None
        mask_override_label: str | None = None
        if mask_option and not selected_labels_and_paths:
            candidate_path = mask_option.get("path")
            if candidate_path and os.path.exists(candidate_path):
                mask_override_path = candidate_path
                mask_override_label = mask_option.get("label")
            else:
                logger.warning("指定的掩码图像不存在或不可访问：%s", candidate_path)

        overlay_preview = _resolve_preview_path("overlay_path", "overlay_original")

        if selected_labels_and_paths:
            labels_and_paths = []
            for index, (detail_label, path, color_suffix) in enumerate(
                selected_labels_and_paths
            ):
                default_label = (
                    channel_signal_labels[index]
                    if index < len(channel_signal_labels)
                    else detail_label
                    or f"{channel_label_prefix}通道图像 {index + 1}"
                )
                label = _resolve_channel_label_from_suffix(
                    channel_label_prefix, color_suffix, default_label
                )
                labels_and_paths.append((label, path))
        else:
            channel_order = [
                ("绿色信号", ("green_path", "green_original")),
                ("红色信号", ("red_path", "red_original")),
                ("蓝色信号", ("blue_path", "blue_original")),
            ]
            labels_and_paths = [
                (
                    f"{channel_label_prefix}{suffix}",
                    _resolve_preview_path(*keys),
                )
                for suffix, keys in channel_order
            ]

            if mask_override_path and len(labels_and_paths) > 1:
                detail_label = mask_override_label or "掩码图像"
                labels_and_paths[1] = (detail_label, mask_override_path)

        if overlay_preview:
            labels_and_paths.append(("叠加合并", overlay_preview))

        valid_items = [
            (label, path)
            for label, path in labels_and_paths
            if path and os.path.exists(path)
        ]

        if valid_items:
            column_count = 2 if len(valid_items) > 2 else len(valid_items)
            column_count = max(column_count, 1)
            row_groups = math.ceil(len(valid_items) / column_count)
            image_table = document.add_table(rows=row_groups * 2, cols=column_count)
            image_table.autofit = True
            image_table.style = None
            _set_table_transparent(image_table)

            for idx, (label, path) in enumerate(valid_items):
                group_index = idx // column_count
                column_index = idx % column_count
                image_row_index = group_index * 2
                label_row_index = image_row_index + 1

                image_row = image_table.rows[image_row_index]
                label_row = image_table.rows[label_row_index]

                image_cell = image_row.cells[column_index]
                image_paragraph = image_cell.paragraphs[0]
                image_run = image_paragraph.add_run()
                image_run.add_picture(path, width=Inches(2.0))
                image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

                label_cell = label_row.cells[column_index]
                _set_cell_text(label_cell, label, bold=True, color=RGBColor(31, 41, 55))
                label_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 清空未使用的单元格，避免残留默认段落
            total_slots = row_groups * column_count
            if total_slots > len(valid_items):
                for empty_idx in range(len(valid_items), total_slots):
                    group_index = empty_idx // column_count
                    column_index = empty_idx % column_count
                    image_row_index = group_index * 2
                    label_row_index = image_row_index + 1
                    image_cell = image_table.rows[image_row_index].cells[column_index]
                    image_cell.text = ""
                    label_cell = image_table.rows[label_row_index].cells[column_index]
                    label_cell.text = ""
        else:
            no_image_paragraph = document.add_paragraph()
            no_image_run = no_image_paragraph.add_run("当前未检测到可用于展示的CTC图像。")
            _apply_run_style(no_image_run, BODY_FONT_SIZE_PT, color=RGBColor(107, 114, 128))
    else:
        no_image_paragraph = document.add_paragraph()
        no_image_run = no_image_paragraph.add_run("当前未检测到可用于展示的CTC图像。")
        _apply_run_style(no_image_run, BODY_FONT_SIZE_PT, color=RGBColor(107, 114, 128))

    description_heading, description_lines = (
        result_description
        if result_description is not None
        else _generate_result_description(metadata, ctc_counts, wbc_counts)
    )
    result_color = RGBColor(0, 0, 0) if doc_names else RGBColor(107, 114, 128)

    result_paragraph = document.add_paragraph()
    result_paragraph.paragraph_format.space_after = Pt(0 if description_lines else 4)
    result_run = result_paragraph.add_run(description_heading)
    _apply_run_style(result_run, BODY_FONT_SIZE_PT, color=result_color)

    for idx, line in enumerate(description_lines):
        line_paragraph = document.add_paragraph()
        line_run = line_paragraph.add_run(line)
        _apply_run_style(line_run, BODY_FONT_SIZE_PT, color=result_color)
        line_paragraph.paragraph_format.space_after = Pt(4 if idx == len(description_lines) - 1 else 0)

    biomarker_value = (metadata.get("癌症标志物") or "").strip()

    notes_value = (metadata.get("备注") or "").strip()
    if notes_value:
        notes_paragraph = document.add_paragraph()
        notes_run = notes_paragraph.add_run(f"备注：{notes_value}")
        _apply_run_style(notes_run, BODY_FONT_SIZE_PT, color=RGBColor(30, 64, 45))

    separator = document.add_paragraph()
    separator_run = separator.add_run("─────────────────────────────────────────────────")
    _apply_run_style(separator_run, BODY_FONT_SIZE_PT, color=RGBColor(75, 85, 99))
    separator.alignment = WD_ALIGN_PARAGRAPH.CENTER


    for index, line in enumerate(DISCLAIMER_LINES):
        disclaimer_paragraph = document.add_paragraph()
        disclaimer_paragraph.paragraph_format.space_before = Pt(4 if index == 0 else 2)
        disclaimer_paragraph.paragraph_format.space_after = Pt(0)
        disclaimer_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        disclaimer_run = disclaimer_paragraph.add_run(line)
        _apply_run_style(disclaimer_run, BODY_FONT_SIZE_PT, color=RGBColor(107, 114, 128))

    document.save(output_path)


def _prepare_workdir(contents: bytes) -> tuple[Path, Path]:
    work_dir = Path(tempfile.mkdtemp(prefix="ctc-analysis-"))
    try:
        archive_path = work_dir / "upload.zip"
        archive_path.write_bytes(contents)

        with zipfile.ZipFile(archive_path) as archive:
            extracted_dir = work_dir / "extracted"
            _safe_extract(archive, extracted_dir)

        dataset_root = _find_dataset_root(work_dir / "extracted")
    except Exception:
        shutil.rmtree(work_dir, ignore_errors=True)
        raise

    return work_dir, dataset_root


async def _prepare_workdir_from_files(files: List[UploadFile]) -> tuple[Path, Path]:
    work_dir = Path(tempfile.mkdtemp(prefix="ctc-analysis-"))
    dataset_root = work_dir / "uploaded"

    try:
        dataset_root.mkdir(parents=True, exist_ok=True)
        root_resolved = dataset_root.resolve()

        for upload in files:
            filename = upload.filename
            if not filename:
                continue

            target_path = root_resolved / filename
            target_path.parent.mkdir(parents=True, exist_ok=True)

            try:
                target_path.resolve().relative_to(root_resolved)
            except ValueError as exc:
                raise HTTPException(status_code=400, detail="上传的文件路径无效") from exc

            content = await upload.read()
            if not content:
                continue
            target_path.write_bytes(content)

        dataset_dir = _find_dataset_root(dataset_root)
    except Exception:
        shutil.rmtree(work_dir, ignore_errors=True)
        raise

    return work_dir, dataset_dir


def _resolve_user_output_directory(path_str: str) -> Path | None:
    """Resolve the user specified output directory and ensure it exists."""

    normalized = path_str.strip()
    if not normalized:
        return None

    candidate = Path(normalized).expanduser()
    try:
        candidate.mkdir(parents=True, exist_ok=True)
    except OSError as exc:  # pragma: no cover - 依赖运行环境
        logger.exception("创建输出文件夹失败：%s", candidate)
        raise HTTPException(status_code=400, detail="输出文件夹路径不可用") from exc

    try:
        resolved = candidate.resolve()
    except OSError:
        resolved = candidate

    if not resolved.is_dir():
        raise HTTPException(status_code=400, detail="输出文件夹路径不是文件夹")

    return resolved


def _resolve_mask_input_directory(path_str: str) -> Path:
    """Resolve the directory that contains external mask images."""

    normalized = path_str.strip()
    if not normalized:
        raise HTTPException(status_code=400, detail="掩码输入文件夹路径不可为空")

    candidate = Path(normalized).expanduser()
    try:
        resolved = candidate.resolve()
    except OSError:
        resolved = candidate

    if not resolved.exists():
        raise HTTPException(status_code=400, detail="掩码输入文件夹不存在")

    if not resolved.is_dir():
        raise HTTPException(status_code=400, detail="掩码输入文件夹不是文件夹")

    return resolved


def _import_mask_images_from_directory(mask_dir: Path, output_dir: Path) -> tuple[list[dict], list[str]]:
    """Import mask images from an external directory into the output directory."""

    warnings: list[str] = []
    mask_files = sorted(
        (
            path
            for path in mask_dir.rglob("*")
            if path.is_file() and path.suffix.lower() in MASK_IMAGE_EXTENSIONS
        ),
        key=lambda item: item.as_posix(),
    )

    if not mask_files:
        warnings.append("掩码输入文件夹中未找到支持的图像文件")
        return [], warnings

    target_root = output_dir / "external_masks"
    options: list[dict] = []

    for source_path in mask_files:
        try:
            relative_source = source_path.relative_to(mask_dir)
        except ValueError:
            relative_source = Path(source_path.name)

        destination = target_root / relative_source
        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source_path, destination)
        except shutil.SameFileError:
            logger.debug("掩码图像已位于目标目录，跳过复制：%s", source_path)
            continue
        except FileNotFoundError:
            warnings.append(f"无法导入掩码图像：{relative_source}")
            logger.warning("掩码图像复制失败（未找到源文件）：%s", source_path)
            continue
        except OSError as exc:
            warnings.append(f"无法导入掩码图像：{relative_source}")
            logger.warning("掩码图像复制失败：%s -> %s (%s)", source_path, destination, exc)
            continue

        encoded_preview = _encode_preview_image(destination)
        if not encoded_preview:
            warnings.append(f"无法读取掩码图像：{relative_source}")
            logger.warning("掩码图像预览生成失败：%s", destination)
            continue

        mime_type, encoded_data = encoded_preview
        try:
            relative_output = destination.relative_to(output_dir)
        except ValueError:
            relative_output = Path(destination.name)

        option_id = f"external::{relative_output.as_posix()}"
        label = relative_source.as_posix() if relative_source.parts else destination.name

        options.append(
            {
                "id": option_id,
                "channel": "external",
                "label": label,
                "relativePath": str(relative_output).replace(os.sep, "/"),
                "mimeType": mime_type or "image/png",
                "data": encoded_data,
            }
        )

    if options:
        logger.info("已导入外部掩码图像 %d 张：%s", len(options), mask_dir)

    return options, warnings


def _sync_output_to_user_directory(
    source_dir: Path,
    user_dir: Path | None,
    *,
    copy_b: bool = True,
) -> None:
    """Copy generated artifacts into the user provided directory."""

    if user_dir is None:
        return

    try:
        user_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:  # pragma: no cover - 依赖运行环境
        logger.exception("无法创建输出文件夹：%s", user_dir)
        raise HTTPException(status_code=400, detail="输出文件夹路径不可用") from exc

    try:
        if source_dir.resolve() == user_dir.resolve():
            return
    except OSError:
        pass

    if copy_b:
        # 查找时间戳命名的子目录而不是固定的"b"目录
        timestamp_dir = None
        for item in source_dir.iterdir():
            if item.is_dir() and len(item.name) == 15 and item.name[8] == '_':  # 格式 YYYYMMDD_HHMMSS
                try:
                    datetime.strptime(item.name, "%Y%m%d_%H%M%S")
                    timestamp_dir = item
                    break
                except ValueError:
                    continue
        
        if timestamp_dir and timestamp_dir.exists():
            target_timestamp_dir = user_dir / timestamp_dir.name
            if target_timestamp_dir.exists():
                shutil.rmtree(target_timestamp_dir, ignore_errors=True)
            try:
                shutil.copytree(timestamp_dir, target_timestamp_dir)
            except OSError as exc:
                logger.exception("复制时间戳文件夹失败：%s -> %s", timestamp_dir, target_timestamp_dir)
                raise HTTPException(status_code=500, detail="同步输出文件夹失败") from exc
        else:
            logger.warning("源目录缺少时间戳文件夹：%s", source_dir)

    report_path = source_dir / "ctc_report.docx"
    if report_path.exists():
        try:
            shutil.copy2(report_path, user_dir / "ctc_report.docx")
        except OSError as exc:
            logger.exception("复制报告文档失败：%s", report_path)
            raise HTTPException(status_code=500, detail="同步输出文件夹失败") from exc

    external_masks_dir = source_dir / "external_masks"
    if external_masks_dir.exists():
        target_external_masks = user_dir / "external_masks"
        if target_external_masks.exists():
            shutil.rmtree(target_external_masks, ignore_errors=True)
        try:
            shutil.copytree(external_masks_dir, target_external_masks)
        except OSError as exc:
            logger.exception("复制外部掩码图像失败：%s", external_masks_dir)
            raise HTTPException(status_code=500, detail="同步输出文件夹失败") from exc


def _prepare_output_directory() -> Path:
    """Create the persistent output directory for generated artifacts."""

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)

    folder_name = timestamp
    target_dir = OUTPUT_ROOT / folder_name
    suffix = 1
    while target_dir.exists():
        target_dir = OUTPUT_ROOT / f"{folder_name}_{suffix}"
        suffix += 1

    target_dir.mkdir(parents=True, exist_ok=False)
    return target_dir


verify_startup_token()

APP_START_TIME = time.time()

def _determine_output_root() -> Path:
    """Return a writable directory for generated artifacts."""
    if getattr(sys, "frozen", False):
        base_dir = Path(sys.executable).resolve().parent
    else:
        base_dir = Path(__file__).resolve().parents[1]

    output_root = base_dir / "var"
    output_root.mkdir(parents=True, exist_ok=True)
    return output_root


OUTPUT_ROOT = _determine_output_root()

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/healthz")
async def heartbeat() -> dict[str, float | str]:
    """简单的心跳检测端点，供前端检测后端状态"""
    uptime_seconds = time.time() - APP_START_TIME
    return {
        "status": "ok",
        "timestamp": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "uptime": round(uptime_seconds, 3),
    }


@app.get("/ctc/sample-number/next")
async def fetch_next_sample_number(
    pet_type: str | None = Query(None, alias="petType")
) -> dict[str, str]:
    normalized_pet_type = _normalize_pet_type(pet_type)

    sample_number_value = _generate_next_sample_number(normalized_pet_type)
    report_number_value = _generate_next_report_number()
    logger.info(
        "通过接口生成样本编号：%s，报告编号：%s（宠物类型=%s）",
        sample_number_value,
        report_number_value,
        normalized_pet_type,
    )
    return {
        "sampleNumber": sample_number_value,
        "reportNumber": report_number_value,
    }


@app.post("/ctc/report")
async def generate_ctc_report(
    file: UploadFile | None = File(None),
    files: List[UploadFile] | None = File(None),
    institution_name: str = Form("", alias="institutionName"),
    report_number: str = Form("", alias="reportNumber"),
    detection_date: str = Form("", alias="detectionDate"),
    sample_number: str = Form("", alias="sampleNumber"),
    sample_volume: str = Form("", alias="sampleVolume"),
    sample_status: str = Form("", alias="sampleStatus"),
    pet_type: str = Form("", alias="petType"),
    cancer_biomarker: str = Form("", alias="cancerBiomarker"),
    department: str = Form("", alias="department"),
    pet_name: str = Form("", alias="petName"),
    owner_name: str = Form("", alias="ownerName"),
    age: str = Form("", alias="age"),
    gender: str = Form("", alias="gender"),
    sample_type: str = Form("", alias="sampleType"),
    medication_intake: str = Form("", alias="medicationIntake"),
    notes: str = Form("", alias="notes"),
    roundness_threshold: float = Form(0.3, alias="roundnessThreshold"),
    preview_only: str = Form("false", alias="previewOnly"),
    output_dir_path: str = Form("", alias="outputDirPath"),
    mask_input_dir_path: str = Form("", alias="maskInputDirPath"),
) -> JSONResponse:
    uploaded_files = files or []
    preview_only_flag = str(preview_only).lower() in {"1", "true", "yes", "on"}
    logger.info(
        "收到CTC报告生成请求：zip文件=%s，多文件数量=%d，圆度阈值=%.2f",
        bool(file and file.filename),
        len(uploaded_files),
        roundness_threshold,
    )

    if (file is None or not file.filename) and not files:
        raise HTTPException(status_code=400, detail="请上传包含影像数据的ZIP文件或文件夹")

    work_dir: Path | None = None
    dataset_root: Path | None = None
    output_dir: Path | None = None
    user_output_dir: Path | None = None
    mask_input_dir: Path | None = None
    try:
        if file is not None and file.filename:
            if not file.filename.lower().endswith(".zip"):
                raise HTTPException(status_code=400, detail="请上传ZIP格式的影像压缩包")

            contents = await file.read()
            if not contents:
                raise HTTPException(status_code=400, detail="上传文件内容为空")

            work_dir, dataset_root = _prepare_workdir(contents)
            logger.info("已从ZIP文件提取数据，工作目录：%s", work_dir)
        elif files:
            work_dir, dataset_root = await _prepare_workdir_from_files(files)
            logger.info("已接收多文件上传，工作目录：%s", work_dir)
        else:
            raise HTTPException(status_code=400, detail="未提供有效的影像数据")

        user_output_dir = _resolve_user_output_directory(output_dir_path)
        if user_output_dir:
            logger.info("指定输出文件夹：%s", user_output_dir)

        output_dir = _prepare_output_directory()
        # 使用时间戳作为目录名替换原来的"b"
        timestamp_dir_name = datetime.now().strftime("%Y%m%d_%H%M%S")
        processing_output_dir = output_dir / timestamp_dir_name
        processing_output_dir.mkdir(parents=True, exist_ok=True)
        logger.info("报告输出目录：%s", output_dir)

        analyzer = CTCAnalyzer(
            str(dataset_root),
            str(processing_output_dir),
            roundness_threshold=roundness_threshold,
        )
        logger.info("开始处理影像数据，数据根目录：%s", dataset_root)
        analyzer.process_all_images()
        logger.info("图像处理完成，生成统计结果：%s", analyzer.results.get("doc_names", []))

        normalized_pet_type_value = _normalize_pet_type(pet_type)
        sanitized_sample_number = (sample_number or "").strip().upper()
        expected_prefix = _resolve_sample_number_prefix(normalized_pet_type_value)
        generated_sample_number = ""
        if sanitized_sample_number:
            if sanitized_sample_number.startswith(expected_prefix):
                numeric_part = sanitized_sample_number[len(expected_prefix):]
                if numeric_part.isdigit() and len(numeric_part) == 7:
                    _sync_sample_number_counter(expected_prefix, int(numeric_part))
                else:
                    sanitized_sample_number = _generate_next_sample_number(normalized_pet_type_value)
                    generated_sample_number = sanitized_sample_number
            else:
                sanitized_sample_number = _generate_next_sample_number(normalized_pet_type_value)
                generated_sample_number = sanitized_sample_number
        else:
            sanitized_sample_number = _generate_next_sample_number(normalized_pet_type_value)
            generated_sample_number = sanitized_sample_number

        normalized_report_number_input = (report_number or "").strip().upper()
        normalized_report_number = normalized_report_number_input
        generated_report_number = ""
        if normalized_report_number_input:
            if normalized_report_number_input.startswith(REPORT_NUMBER_PREFIX):
                numeric_part = normalized_report_number_input[len(REPORT_NUMBER_PREFIX):]
                if numeric_part.isdigit() and len(numeric_part) == SAMPLE_NUMBER_DIGITS:
                    _sync_report_number_counter(int(numeric_part))
                    normalized_report_number = f"{REPORT_NUMBER_PREFIX}{numeric_part}"
                else:
                    normalized_report_number = _generate_next_report_number()
                    generated_report_number = normalized_report_number
            else:
                normalized_report_number = _generate_next_report_number()
                generated_report_number = normalized_report_number
        else:
            normalized_report_number = _generate_next_report_number()
            generated_report_number = normalized_report_number

        metadata = _create_metadata_entries(
            institution_name,
            normalized_report_number,
            detection_date,
            sanitized_sample_number,
            sample_volume,
            sample_status,
            normalized_pet_type_value or pet_type,
            cancer_biomarker,
            department,
            pet_name,
            owner_name,
            age,
            gender,
            sample_type,
            medication_intake,
            notes,
        )
        logger.info(
            "报告元数据：%s",
            {key: metadata[key] for key in metadata},
        )
        auto_generated_sample_number = bool(generated_sample_number)
        auto_generated_report_number = bool(generated_report_number)

        report_path = output_dir / "ctc_report.docx"
        channel_summary_texts = _format_channel_summary_texts(analyzer)

        doc_names = analyzer.results.get("doc_names", [])
        ctc_counts = analyzer.results.get("green_single_channel", [])
        wbc_counts = analyzer.results.get("white_single_channel", [])
        channel_stats = [
            {"channel": name, "ctc": ctc, "wbc": wbc}
            for name, ctc, wbc in zip(doc_names, ctc_counts, wbc_counts)
        ]

        detection_result_rows = _format_detection_result_rows(
            doc_names,
            ctc_counts,
            wbc_counts,
        )

        total_ctc = sum(ctc_counts)
        total_wbc = sum(wbc_counts)

        result_description = _generate_result_description(metadata, ctc_counts, wbc_counts)
        selection_text = ""

        mask_options_payload: list[dict] = []
        mask_paths_map = analyzer.results.get("mask_paths", {}) or {}
        seen_relative_paths: set[str] = set()

        def _append_mask_option(
            path_obj: Path,
            channel_label: str,
            label: str,
            *,
            option_id: str | None = None,
        ) -> None:
            if not path_obj.exists():
                return
            encoded_preview = _encode_preview_image(path_obj)
            if not encoded_preview:
                return
            mime_type, encoded_data = encoded_preview
            try:
                relative_path_obj = path_obj.relative_to(output_dir)
            except ValueError:
                relative_path_obj = Path(path_obj.name)
            normalized_relative = str(relative_path_obj).replace(os.sep, "/")
            if normalized_relative in seen_relative_paths:
                return
            seen_relative_paths.add(normalized_relative)
            mask_options_payload.append(
                {
                    "id": option_id or f"{channel_label}::{normalized_relative}",
                    "channel": channel_label,
                    "label": label,
                    "relativePath": normalized_relative,
                    "mimeType": mime_type or "image/png",
                    "data": encoded_data,
                }
            )

        for channel, mask_paths in mask_paths_map.items():
            channel_str = str(channel)
            directories_for_channel: set[Path] = set()
            for index, mask_path in enumerate(mask_paths):
                mask_path_obj = Path(mask_path)
                if not mask_path_obj.exists():
                    continue
                directories_for_channel.add(mask_path_obj.parent)
                option_label = (
                    f"{channel_str}通道掩码"
                    if len(mask_paths) == 1
                    else f"{channel_str}通道掩码 {index + 1}"
                )
                option_id = f"{channel_str}-{mask_path_obj.stem}-{index}"
                _append_mask_option(
                    mask_path_obj,
                    channel_str,
                    option_label,
                    option_id=option_id,
                )

            for directory in sorted(
                directories_for_channel, key=lambda item: item.as_posix()
            ):
                for candidate in sorted(
                    directory.rglob("*"), key=lambda item: item.as_posix()
                ):
                    if (
                        candidate.is_file()
                        and candidate.suffix.lower() in MASK_IMAGE_EXTENSIONS
                    ):
                        label = f"{channel_str}通道图像：{candidate.name}"
                        _append_mask_option(candidate, channel_str, label)

        if not mask_options_payload:
            fallback_root = output_dir / "b"
            if fallback_root.exists():
                for candidate in sorted(
                    fallback_root.rglob("*"), key=lambda item: item.as_posix()
                ):
                    if (
                        candidate.is_file()
                        and candidate.suffix.lower() in MASK_IMAGE_EXTENSIONS
                    ):
                        _append_mask_option(candidate, "all", candidate.name)

        warnings: list[str] = []

        mask_input_dir_value = mask_input_dir_path.strip()
        if mask_input_dir_value:
            mask_input_dir = _resolve_mask_input_directory(mask_input_dir_value)
            logger.info("指定掩码输入文件夹：%s", mask_input_dir)
            external_mask_options, external_warnings = _import_mask_images_from_directory(
                mask_input_dir, output_dir
            )
            mask_options_payload.extend(external_mask_options)
            warnings.extend(external_warnings)
        else:
            mask_input_dir = None

        _sync_output_to_user_directory(output_dir, user_output_dir)

        ctc_image_sets = analyzer.results.get("ctc_image_sets", [])
        image_set_payload: dict | None = None
        if ctc_image_sets:
            first_set = ctc_image_sets[0]
            channel_label_prefix = _format_channel_label_prefix(
                first_set.get("sub_folder")
            )

            def _resolve_image(primary_key: str, *fallback_keys: str) -> Path | None:
                keys = (primary_key,) + fallback_keys
                for key in keys:
                    candidate = first_set.get(key)
                    if candidate:
                        candidate_path = Path(candidate)
                        if candidate_path.exists():
                            return candidate_path
                return None

            images_payload = []
            for suffix, keys in (
                ("绿色信号", ("green_path", "green_original")),
                ("红色信号", ("red_path", "red_original")),
                ("蓝色信号", ("blue_path", "blue_original")),
                ("叠加通道", ("overlay_path", "overlay_original")),
            ):
                image_path = _resolve_image(*keys)
                if not image_path:
                    continue
                encoded_preview = _encode_preview_image(image_path)
                if not encoded_preview:
                    continue
                mime_type, encoded_data = encoded_preview
                label = (
                    f"{channel_label_prefix}{suffix}"
                    if suffix != "叠加通道"
                    else suffix
                )
                images_payload.append(
                    {
                        "label": label,
                        "mimeType": mime_type or "image/png",
                        "data": encoded_data,
                    }
                )

            if images_payload:
                image_set_payload = {"items": images_payload}

        heading_text, description_lines = result_description
        result_text_parts = [heading_text, *(description_lines or [])]
        result_text = "\n".join(part for part in result_text_parts if part)

        notes_value = (metadata.get("备注") or "").strip()
        remark_lines: list[str] = []
        if notes_value:
            remark_lines.append(f"备注：{notes_value}")

        remark_text = "\n".join(remark_lines)

        metadata_items = [
            {"label": label, "value": value}
            for label, value in metadata.items()
        ]

        report_token = output_dir.name
        state_payload = {
            "metadata": {key: metadata[key] for key in metadata},
            "results": analyzer.results,
            "channelSummaryTexts": channel_summary_texts,
            "selectionText": selection_text,
            "resultText": result_text,
            "detectionResultRows": detection_result_rows,
            "remarkText": remark_text,
            "channelLabelPrefix": channel_label_prefix,
            "maskOptions": [
                {
                    "id": option["id"],
                    "channel": option["channel"],
                    "label": option["label"],
                    "relativePath": option["relativePath"],
                }
                for option in mask_options_payload
            ],
            "createdAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
            "userOutputDirectory": str(user_output_dir) if user_output_dir else "",
            "maskInputDirectory": str(mask_input_dir) if mask_input_dir else "",
            "generatedSampleNumber": sanitized_sample_number,
            "autoGeneratedSampleNumber": auto_generated_sample_number,
            "generatedReportNumber": normalized_report_number,
            "autoGeneratedReportNumber": auto_generated_report_number,
        }
        (output_dir / "report_state.json").write_text(
            json.dumps(state_payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        response_payload = {
            "fileName": "ctc_report.docx",
            "generatedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
            "metadata": metadata_items,
            "channels": channel_stats,
            "totals": {
                "totalCtc": total_ctc,
                "totalWbc": total_wbc,
            },
            "resultText": result_text,
            "detectionResultRows": detection_result_rows,
            "remarkText": remark_text,
            "selectionText": selection_text,
            "hasCtcImages": bool(image_set_payload),
            "imageSet": image_set_payload,
            "channelSummaryTexts": channel_summary_texts,
            "channelLabelPrefix": channel_label_prefix,
            "reportToken": report_token,
            "maskOptions": mask_options_payload,
            "userOutputDirectory": str(user_output_dir) if user_output_dir else "",
            "maskInputDirectory": str(mask_input_dir) if mask_input_dir else "",
            "warnings": warnings,
            "generatedSampleNumber": sanitized_sample_number,
            "autoGeneratedSampleNumber": auto_generated_sample_number,
            "generatedReportNumber": normalized_report_number,
            "autoGeneratedReportNumber": auto_generated_report_number,
        }

        if not preview_only_flag:
            _build_report_document(
                analyzer,
                metadata,
                report_path,
                channel_summary_texts=channel_summary_texts,
                detection_result_rows=detection_result_rows,
                result_description=result_description,
            )
            logger.info("报告已生成：%s", report_path)
            _sync_output_to_user_directory(output_dir, user_output_dir, copy_b=False)
            report_bytes = report_path.read_bytes()
            encoded_report = base64.b64encode(report_bytes).decode("ascii")
            response_payload["fileContent"] = encoded_report

        shutil.rmtree(work_dir, ignore_errors=True)
        return JSONResponse(content=response_payload)
    except zipfile.BadZipFile as exc:
        if work_dir is not None:
            shutil.rmtree(work_dir, ignore_errors=True)
        if output_dir is not None and output_dir.exists() and not any(output_dir.iterdir()):
            output_dir.rmdir()
        raise HTTPException(status_code=400, detail="压缩包文件损坏或格式错误") from exc
    except HTTPException:
        if work_dir is not None:
            shutil.rmtree(work_dir, ignore_errors=True)
        if output_dir is not None and output_dir.exists() and not any(output_dir.iterdir()):
            output_dir.rmdir()
        raise
    except Exception as exc:
        logger.exception("生成CTC报告时出现未预期错误")
        if work_dir is not None:
            shutil.rmtree(work_dir, ignore_errors=True)
        if output_dir is not None and output_dir.exists() and not any(output_dir.iterdir()):
            output_dir.rmdir()
        raise HTTPException(status_code=500, detail="生成报告时发生未知错误") from exc
    finally:
        if file is not None:
            await file.close()
        if files:
            for upload in files:
                await upload.close()


@app.post("/ctc/report/export")
async def export_ctc_report(
    report_token: str = Form(..., alias="reportToken"),
    mask_option_id: str = Form("", alias="maskOptionId"),
    mask_option_ids: str = Form("", alias="maskOptionIds"),
) -> JSONResponse:
    logger.info(
        "收到报告导出请求：report_token=%s, mask_option_id=%s, mask_option_ids=%s",
        report_token,
        mask_option_id,
        mask_option_ids,
    )

    if not report_token:
        raise HTTPException(status_code=400, detail="缺少报告标识")

    output_root = OUTPUT_ROOT.resolve()
    output_dir = (output_root / report_token).resolve()
    try:
        output_dir.relative_to(output_root)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="报告标识无效") from exc

    if not output_dir.exists():
        raise HTTPException(status_code=404, detail="报告数据不存在")

    state_path = output_dir / "report_state.json"
    if not state_path.exists():
        raise HTTPException(status_code=400, detail="报告状态文件不存在")

    try:
        state_data = json.loads(state_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        logger.exception("报告状态文件解析失败：%s", state_path)
        raise HTTPException(status_code=500, detail="报告状态文件已损坏") from exc

    metadata_dict = state_data.get("metadata") or {}
    metadata = OrderedDict((key, metadata_dict[key]) for key in metadata_dict)
    results = state_data.get("results") or {}
    doc_names_state = results.get("doc_names") or []
    ctc_counts_state = results.get("green_single_channel") or []
    wbc_counts_state = results.get("white_single_channel") or []
    channel_summary_texts = state_data.get("channelSummaryTexts") or []
    mask_options_state = state_data.get("maskOptions") or []
    user_output_dir_str = state_data.get("userOutputDirectory") or ""
    user_output_dir = (
        _resolve_user_output_directory(user_output_dir_str)
        if user_output_dir_str
        else None
    )

    detection_result_rows_state = state_data.get("detectionResultRows")
    detection_result_rows: list[list[str]] | None = None
    if isinstance(detection_result_rows_state, list):
        parsed_rows: list[list[str]] = []
        for row in detection_result_rows_state:
            if not isinstance(row, list):
                parsed_rows = []
                break
            parsed_row: list[str] = []
            for item in row:
                if isinstance(item, str):
                    parsed_row.append(item)
            if parsed_row:
                parsed_rows.append(parsed_row)
        if parsed_rows:
            detection_result_rows = parsed_rows

    if detection_result_rows is None:
        detection_result_rows = _format_detection_result_rows(
            doc_names_state,
            ctc_counts_state,
            wbc_counts_state,
        )

    selected_mask_option_ids: list[str] = []
    if mask_option_ids:
        try:
            parsed_ids = json.loads(mask_option_ids)
        except json.JSONDecodeError as exc:
            logger.warning("掩码图像选择解析失败：%s", mask_option_ids)
            raise HTTPException(status_code=400, detail="掩码图像选择格式无效") from exc

        if not isinstance(parsed_ids, list):
            raise HTTPException(status_code=400, detail="掩码图像选择格式无效")

        for item in parsed_ids:
            if not isinstance(item, str):
                raise HTTPException(status_code=400, detail="掩码图像选择格式无效")
            normalized = item.strip()
            if normalized and normalized not in selected_mask_option_ids:
                selected_mask_option_ids.append(normalized)

    elif mask_option_id:
        normalized = mask_option_id.strip()
        if normalized:
            selected_mask_option_ids.append(normalized)

    required_mask_count = 3
    if selected_mask_option_ids and len(selected_mask_option_ids) != required_mask_count:
        raise HTTPException(
            status_code=400,
            detail=f"必须选择{required_mask_count}张掩码图像",
        )

    mask_option_payloads: list[dict] = []
    if selected_mask_option_ids:
        for option_id in selected_mask_option_ids:
            option = next((item for item in mask_options_state if item.get("id") == option_id), None)
            if option is None:
                raise HTTPException(status_code=400, detail="所选掩码图像不存在")

            relative_path = option.get("relativePath")
            if not relative_path:
                raise HTTPException(status_code=400, detail="所选掩码图像路径无效")

            mask_path = (output_dir / relative_path).resolve()
            try:
                mask_path.relative_to(output_dir)
            except ValueError as exc:
                raise HTTPException(status_code=400, detail="掩码图像路径无效") from exc

            if not mask_path.exists():
                raise HTTPException(status_code=400, detail="所选掩码图像文件不存在")

            mask_option_payloads.append(
                {
                    "path": str(mask_path),
                    "label": option.get("label") or f"{option.get('channel', '')}通道掩码",
                }
            )

    result_description = _generate_result_description(
        metadata,
        ctc_counts_state,
        wbc_counts_state,
    )

    analyzer_stub = SimpleNamespace(results=results)
    report_path = output_dir / "ctc_report.docx"

    _build_report_document(
        analyzer_stub,  # type: ignore[arg-type]
        metadata,
        report_path,
        channel_summary_texts=channel_summary_texts,
        mask_options=mask_option_payloads or None,
        detection_result_rows=detection_result_rows,
        result_description=result_description,
    )
    logger.info("报告文档已生成：%s", report_path)

    _sync_output_to_user_directory(output_dir, user_output_dir, copy_b=False)

    report_bytes = report_path.read_bytes()
    encoded_report = base64.b64encode(report_bytes).decode("ascii")

    return JSONResponse(
        content={
            "fileName": "ctc_report.docx",
            "fileContent": encoded_report,
            "generatedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
            "reportToken": report_token,
            "userOutputDirectory": str(user_output_dir) if user_output_dir else "",
        }
    )


if __name__ == "__main__":
    import uvicorn

    port_str = os.environ.get("FASTAPI_PORT", "15000")
    try:
        port = int(port_str)
    except ValueError:
        logger.warning("警告：FASTAPI_PORT 设置无效（%s），使用默认端口 15000", port_str)
        port = 15000

    uvicorn.run(app, host="0.0.0.0", port=port)
